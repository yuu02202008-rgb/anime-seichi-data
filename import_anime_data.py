import json
import re
from pathlib import Path
import openpyxl

SOURCE = Path('/Users/yu/Downloads/アニメ聖地 (7).xlsx')
TARGET = Path('data.js')

existing_places = []
if TARGET.exists():
    match = re.search(r'window\.places = (\[.*?\]);\nwindow\.workInfo', TARGET.read_text(encoding='utf-8'), re.S)
    if match:
        existing_places = json.loads(match.group(1))
existing_by_pair = {}
for place in existing_places:
    existing_by_pair.setdefault((str(place.get('work', '')).strip(), str(place.get('name', '')).strip()), place)

workbook = openpyxl.load_workbook(SOURCE, data_only=True)
data_sheet = workbook['データ']
service_sheet = workbook['類似サービス']

def text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()

headers = [text(data_sheet.cell(4, column).value) for column in range(1, 42)]
work_columns = [
    '連載開始時期', '連載終了時期', 'アニメ初放送', '原作出版社', '連載誌', 'アニメーション制作会社',
    '作品名', '作品名カナ', '作品名略称', 'ジャンル1', 'ジャンル2', '作者', '作者カナ',
    '主人公', '主人公カナ', '主要人物1', '主要人物1カナ', '主要人物2', '主要人物2カナ',
    '主要人物3', '主要人物3カナ', 'ストーリー', '漫画', 'アニメ', '映画', 'Youtube', 'その他'
]
works = {}
places = []
seen_pairs = set()
used_ids = set()
colors = ['orange', 'blue', 'yellow', 'green', 'red', 'purple']
verified_extras = [
    {
        'id': 'verified-chichibu-station-anohana', 'name': '秩父駅', 'prefecture': '埼玉県', 'city': '秩父市',
        'category': '駅', 'work': 'あの日見た花の名前を僕達はまだ知らない。',
        'scene': '秩父市公式の聖地巡礼モデルコースに「あの花」ゆかりの移動地点として掲載',
        'episode': '作品ゆかりの場所', 'confidence': 'A', 'checkedAt': '2026-08-28',
        'coordinates': '35.99960, 139.08649', 'address': '埼玉県秩父市宮側町1-8', 'nearestStation': '秩父駅',
        'visit': '条件付き', 'visitConditions': '駅構内・ホームへ入る場合は有効な乗車券または入場券が必要です。',
        'mapUrl': 'https://www.google.com/maps/search/?api=1&query=秩父駅', 'privacyProtected': False,
        'sourceUrl': 'https://navi.city.chichibu.lg.jp/travel/ja/course/2018/11/925/'
    },
    {
        'id': 'verified-kasukabe-station-crayon', 'name': '春日部駅', 'prefecture': '埼玉県', 'city': '春日部市',
        'category': '駅', 'work': 'クレヨンしんちゃん',
        'scene': '春日部市公式資料で、漫画・アニメでおなじみの場所として紹介されている駅',
        'episode': '作品ゆかりの場所', 'confidence': 'A', 'checkedAt': '2026-08-28',
        'coordinates': '35.97999, 139.75239', 'address': '埼玉県春日部市粕壁1-10-1', 'nearestStation': '春日部駅',
        'visit': '条件付き', 'visitConditions': '駅構内・ホームへ入る場合は有効な乗車券または入場券が必要です。高架化工事により構内状況が変わる場合があります。',
        'mapUrl': 'https://www.google.com/maps/search/?api=1&query=春日部駅', 'privacyProtected': False,
        'sourceUrl': 'https://www.city.kasukabe.lg.jp/material/files/group/6/walkP11.pdf'
    },
    {
        'id': 'verified-kamei-arena-haikyu', 'name': 'カメイアリーナ仙台（仙台市体育館）', 'prefecture': '宮城県', 'city': '仙台市太白区',
        'category': '体育館', 'work': 'ハイキュー!!',
        'scene': '仙台市が「ハイキュー!!」仙台観光特使記念モニュメントを設置した作品ゆかりの施設',
        'episode': '作品ゆかりの場所', 'confidence': 'A', 'checkedAt': '2026-08-28',
        'coordinates': '38.21713, 140.86960', 'address': '宮城県仙台市太白区富沢1-4-1', 'nearestStation': '富沢駅',
        'visit': '条件付き', 'visitConditions': '大規模改修のため2027年3月15日予定まで全館休館中です。記念モニュメントは隣接する富沢公園内へ一時移設されています。',
        'mapUrl': 'https://www.google.com/maps/search/?api=1&query=カメイアリーナ仙台', 'privacyProtected': False,
        'sourceUrl': 'https://www.city.sendai.jp/contentsoushutsu/animation_manga/tokushi/haikyu819.html'
    }
]

for row_number in range(5, data_sheet.max_row + 1):
    values = [data_sheet.cell(row_number, column).value for column in range(1, 42)]
    if not values[7] or not values[28]:
        continue
    work_name = text(values[7])
    if work_name not in works:
        works[work_name] = {key: text(value) for key, value in zip(work_columns, values[1:28])}

    spot = text(values[28])
    pair = (work_name, spot)
    if pair in seen_pairs:
        continue
    seen_pairs.add(pair)
    existing = existing_by_pair.get(pair, {})
    proposed_id = existing.get('id') or f'place-{int(values[0]) if values[0] else row_number}'
    if proposed_id in used_ids:
        proposed_id = f'place-new-{row_number}'
    used_ids.add(proposed_id)
    scene = text(values[30])
    address = text(values[34])
    raw_visit = text(values[29])
    visit = raw_visit
    visit_conditions = ''
    if raw_visit == '施設利用要':
        visit = '条件付き'
        visit_conditions = '施設の営業時間・料金・利用ルールを事前に確認し、施設利用者として訪問してください。'
    elif raw_visit == '外観のみ・閉館':
        visit = '外観のみ'
        visit_conditions = '施設は閉館しています。敷地へ立ち入らず、周辺の通行や近隣住民に配慮してください。'
    elif raw_visit == '外観のみ・閉店':
        visit = '外観のみ'
        visit_conditions = '店舗は閉店しています。敷地へ立ち入らず、周辺の通行や近隣施設に配慮してください。'
    protected = any(word in f'{spot} {scene} {address}' for word in ['個人宅', '住宅地', '社宅'])
    place = {
        'id': proposed_id,
        'name': spot,
        'prefecture': text(values[31]),
        'city': text(values[32]),
        'category': '聖地スポット',
        'work': work_name,
        'scene': scene,
        'episode': '作品内のシーン',
        'confidence': 'B',
        'checkedAt': '2026-08-24',
        'coordinates': '非公開（住宅地保護）' if protected else f'{text(values[35])}, {text(values[36])}',
        'address': f'{text(values[31])}{text(values[32])}（詳細非公開）' if protected else address,
        'nearestStation': text(values[33]),
        'visit': visit,
        'mapUrl': '' if protected else text(values[37]),
        'privacyProtected': protected,
        'color': existing.get('color') or colors[(len(places)) % len(colors)]
    }
    if visit_conditions:
        place['visitConditions'] = visit_conditions
    places.append(place)

for extra in verified_extras:
    pair = (extra['work'], extra['name'])
    if pair in seen_pairs:
        continue
    seen_pairs.add(pair)
    extra['color'] = existing_by_pair.get(pair, {}).get('color') or colors[len(places) % len(colors)]
    places.append(extra)

payload = (
    '// このファイルは work/import_anime_data.py で生成されます。\n'
    '// 出典：ユーザー提供「アニメ聖地 (7).xlsx」＋公式情報で独自確認した追加地点（作品名＋聖地名の完全一致を除外）\n'
    f'window.places = {json.dumps(places, ensure_ascii=False, separators=(",", ":"))};\n'
    f'window.workInfo = {json.dumps(works, ensure_ascii=False, separators=(",", ":"))};\n'
)
TARGET.write_text(payload, encoding='utf-8')
print(f'Imported {len(places)} places and {len(works)} works.')
